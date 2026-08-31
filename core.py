import os
import json
import io
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, field_validator

from pypdf import PdfReader
from docx import Document

load_dotenv()
_api_key = os.getenv("GROQ_API_KEY")

if not _api_key:
    raise ValueError("GROQ_API_KEY not found. Add it to your .env file.")

client = Groq(api_key=_api_key)
MODEL = "openai/gpt-oss-120b"


# ---------- Schemas ----------

class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]

    @field_validator(
        "required_skills", "preferred_skills", "education_requirements", "responsibilities",
        mode="before",
    )
    @classmethod
    def _null_to_empty_list(cls, v):
        return v if v is not None else []


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

    @field_validator("skills_used", mode="before")
    @classmethod
    def _null_to_empty_list(cls, v):
        return v if v is not None else []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []

    @field_validator("skills", "education", "projects", "certifications", "experiences", mode="before")
    @classmethod
    def _null_to_empty_list(cls, v):
        return v if v is not None else []


class MatchResult(BaseModel):
    candidate_name: str | None = None
    score: float
    matching_skills: list[str] = []
    missing_skills: list[str] = []
    experience_requirement_met: bool | None = None
    verdict: str
    feedback: str | None = None  # only populated in individual mode

    @field_validator("matching_skills", "missing_skills", mode="before")
    @classmethod
    def _null_to_empty_list(cls, v):
        return v if v is not None else []


# ---------- File reading ----------

def read_pdf(file_obj) -> str:
    reader = PdfReader(file_obj)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx(file_obj) -> str:
    document = Document(file_obj)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(uploaded_file) -> str | None:
    """
    Accepts a Streamlit UploadedFile (has .name and behaves like a file object).
    """
    name = uploaded_file.name.lower()
    file_bytes = io.BytesIO(uploaded_file.getvalue())

    if name.endswith(".pdf"):
        return read_pdf(file_bytes)
    elif name.endswith(".docx"):
        return read_docx(file_bytes)
    else:
        return None


# ---------- LLM calls ----------

def extract_job_description(jd_text: str) -> JobD:
    schema = JobD.model_json_schema()
    system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job descriptions and extract structured information from them.

Return ONLY valid JSON matching this schema:

{schema}

IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""
    user_prompt = f"Analyze the following job description:\n\n{jd_text}"

    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return JobD(**data)


def parse_resume(resume_text: str) -> Resume:
    schema = Resume.model_json_schema()
    system_prompt = f"""
You are an expert resume parser.

Extract information from the resume based on its meaning,
not only based on exact section headings.

Different resumes may use different headings.
For example: Experience, Professional Experience, Work History, Employment, Internships.
These may all contain relevant experience.

Skills may also appear in the skills section, work experience,
internships or projects.

Return ONLY valid JSON matching this schema:

{schema}

Important rules:
1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experiences.
5. Extract skills mentioned across the entire resume.
"""
    user_prompt = f"Parse the following resume:\n\n{resume_text}"

    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return Resume(**data)


def score_resume(job: JobD, resume: Resume) -> MatchResult:
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
You are an HR recruiter.

Compare the candidate's resume with the job description.

JOB DESCRIPTION:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}

Return JSON matching this schema:

{match_schema}

Fill in:
- candidate_name: the candidate's name
- matching_skills: skills from the resume that match the JD's required or preferred skills
- missing_skills: important skills the JD asks for that are absent from the resume
- experience_requirement_met: true/false based on the JD's minimum experience vs the candidate's experience
- score: overall match percentage from 0 to 100
- verdict: a short (1-2 sentence) final verdict

Keep it concise and specific - use actual skill/experience names from the resume and JD, not vague summaries.
Leave "feedback" as null - it will be filled separately.
"""
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    data.setdefault("feedback", None)
    return MatchResult(**data)


def generate_feedback(job: JobD, resume: Resume, match: MatchResult) -> str:
    """
    Only used in individual (self-assessment) mode.
    Gives the candidate concrete advice on improving their resume for this JD.
    """
    prompt = f"""
You are a career coach helping a candidate improve their resume for a specific job.

JOB DESCRIPTION:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}

MATCH ANALYSIS:
{match.model_dump_json(indent=2)}

Write specific, actionable feedback the candidate can apply directly to their resume.
Do NOT give generic advice like "tailor your resume" or "highlight relevant skills" -
every point must name an exact skill, phrase, or line from their resume or the JD.

For each point, use one of these two formats:

- **Add:** "<the exact skill/keyword/phrase>" to your <specific section, e.g. skills
  section or the X project bullet> - because the JD requires it and your resume
  doesn't mention it. (Only suggest adding something the candidate plausibly has
  based on their resume - never invent skills for them.)

- **Change:** "<the current resume line, quoted or closely paraphrased>" -> "<the
  improved version>" - explain in one short clause why the new version matches the
  JD better (e.g. uses the JD's terminology, quantifies impact, surfaces a required skill).

Give 4-6 such points, ordered by impact (most important first). If experience
requirements aren't met, add one final point on how to frame existing experience
to close that gap honestly - do not suggest fabricating experience.
"""
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
    )
    return response.choices[0].message.content